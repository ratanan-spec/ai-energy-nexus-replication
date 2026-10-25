"""Normalize mixed Excel/text dates in the retained Investing.com workbook."""
from datetime import date, datetime
from pathlib import Path
import csv
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "raw" / "energy_markets.xlsx"
OUTPUT = ROOT / "data" / "raw" / "vendor_prices_normalized.csv"
START, END = date(2010, 11, 4), date(2025, 12, 31)
PAIRS = {"NG": (2, 3), "HG": (4, 5), "URA": (6, 7),
         "ICLN": (8, 9), "SOX": (10, 11), "VIX": (12, 13)}


def candidates(value):
    if value is None:
        return []
    if isinstance(value, datetime):
        observed = value.date()
        out = [observed]
        if observed.month <= 12 and observed.day <= 12:
            try:
                out.append(date(observed.year, observed.day, observed.month))
            except ValueError:
                pass
        return list(dict.fromkeys(out))
    text = str(value).strip()
    out = []
    for fmt in ("%m/%d/%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            out.append(datetime.strptime(text, fmt).date())
        except ValueError:
            pass
    return list(dict.fromkeys(out))


def main():
    ws = load_workbook(SOURCE, data_only=True).active
    series = {}
    for name, (date_col, price_col) in PAIRS.items():
        previous, values = None, []
        for row in range(2, ws.max_row + 1):
            raw_date, raw_price = ws.cell(row, date_col).value, ws.cell(row, price_col).value
            if raw_date is None or raw_price is None:
                continue
            valid = [d for d in candidates(raw_date)
                     if START <= d <= END and (previous is None or d > previous)]
            if not valid:
                continue
            chosen = min(valid, key=lambda d: (d - (previous or START)).days)
            values.append((chosen, float(raw_price)))
            previous = chosen
        series[name] = dict(values)
        print(name, len(values), values[0][0], values[-1][0])

    common = sorted(set.intersection(*(set(v) for v in series.values())))
    print("COMMON", len(common), common[0], common[-1])
    if len(common) != 3810 or common[0] != date(2010, 11, 8) or common[-1] != END:
        raise RuntimeError("Unexpected common-date intersection; inspect vendor revisions.")

    all_dates = sorted(set.union(*(set(v) for v in series.values())))
    with OUTPUT.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Date", *PAIRS])
        for day in all_dates:
            writer.writerow([day.isoformat(), *(series[name].get(day, "") for name in PAIRS)])


if __name__ == "__main__":
    main()
